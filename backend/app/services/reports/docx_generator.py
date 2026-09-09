import io
from typing import Any, Dict
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


class DocxReportGenerator:
    """Generates official statutory Microsoft Word (.docx) compliance reports using python-docx."""

    @staticmethod
    def generate_docx_bytes(context: Dict[str, Any]) -> bytes:
        doc = Document()

        # Set standard margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Set default font
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Arial"
        font.size = Pt(9.5)
        font.color.rgb = RGBColor(0x18, 0x1C, 0x21)

        # Title Block
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run("PACKCHECK LEGAL METROLOGY COMPLIANCE REPORT\n")
        title_run.font.size = Pt(15)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x00, 0x51, 0xD5)

        subtitle_run = title_para.add_run("Department of Legal Metrology • Legal Metrology (Packaged Commodities) Rules, 2011")
        subtitle_run.font.size = Pt(9.5)
        subtitle_run.font.color.rgb = RGBColor(0x42, 0x46, 0x55)

        doc.add_paragraph()

        # 1. Metadata Section
        meta_heading = doc.add_heading(level=2)
        meta_run = meta_heading.add_run("1. Inspection & Packaged Commodity Metadata")
        meta_run.font.size = Pt(12)
        meta_run.font.bold = True
        meta_run.font.color.rgb = RGBColor(0x00, 0x51, 0xD5)

        data = [
            ("Report ID:", str(context.get("report_id", "N/A"))),
            ("Scan ID:", str(context.get("scan_id", "N/A"))),
            ("Inspection Date:", str(context.get("generated_at", "N/A"))),
            ("Enforcement Officer:", f"{context.get('inspector_name', 'N/A')} ({context.get('inspector_role', 'Inspector')})"),
            ("Product & Brand:", f"{context.get('product_name', 'N/A')} - {context.get('product_brand', 'N/A')} ({context.get('product_category', 'General')})"),
            ("Manufacturer / Origin:", f"{context.get('manufacturer_name', 'Declared on label')} | Origin: {context.get('country_of_origin', 'India')}"),
        ]

        if context.get("batch_number"):
            data.append(("Batch / Lot Number:", f"{context.get('batch_number')} (Source: {str(context.get('batch_number_source', 'manual')).capitalize()})"))
            b_st = str(context.get("batch_status", "normal")).replace("_", " ").upper()
            rel_cnt = context.get("related_batch_inspections", 1)
            att_cnt = context.get("batch_inspections_requiring_attention", 0)
            data.append(("Batch Investigation Status:", f"{b_st} ({rel_cnt} inspection(s) recorded in batch / {att_cnt} requiring attention)"))

        meta_table = doc.add_table(rows=len(data), cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, (label, val) in enumerate(data):
            cell_0 = meta_table.cell(i, 0)
            cell_1 = meta_table.cell(i, 1)
            cell_0.text = label
            cell_0.paragraphs[0].runs[0].font.bold = True
            cell_0.paragraphs[0].runs[0].font.size = Pt(9)
            cell_1.text = val
            cell_1.paragraphs[0].runs[0].font.size = Pt(9)

        if context.get("batch_number"):
            note_p = doc.add_paragraph()
            note_run = note_p.add_run(f"Batch Traceability Notice: {context.get('batch_legal_note', '')}")
            note_run.font.size = Pt(8)
            note_run.font.italic = True
            note_run.font.color.rgb = RGBColor(0x55, 0x58, 0x69)

        doc.add_paragraph()

        # Overall Status Banner
        status_para = doc.add_paragraph()
        status_label = context.get("overall_label") or str(context.get("overall_status", "N/A")).upper()
        status_run = status_para.add_run(f"Overall Compliance Result: {status_label}\n")
        status_run.font.bold = True
        status_run.font.size = Pt(11.5)
        
        st = context.get("overall_status")
        if st == "compliant":
            status_run.font.color.rgb = RGBColor(0x00, 0x6B, 0x5C)
        elif st == "needs_review":
            status_run.font.color.rgb = RGBColor(0xB2, 0x5E, 0x00)
        else:
            status_run.font.color.rgb = RGBColor(0xBA, 0x1A, 0x1A)

        # 2. Declarations Table
        decl_heading = doc.add_heading(level=2)
        decl_run = decl_heading.add_run("2. Mandatory Declaration Audit (LMPC Rules, 2011)")
        decl_run.font.size = Pt(12)
        decl_run.font.bold = True
        decl_run.font.color.rgb = RGBColor(0x00, 0x51, 0xD5)

        declarations = context.get("declarations", [])
        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        hdr_cells = table.rows[0].cells
        headers = ["Declaration Field", "Extracted Value & Evidence", "Detection Rate & Threshold", "Workflow & Decision", "Rule Reference & Reason"]
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
            hdr_cells[idx].paragraphs[0].runs[0].font.size = Pt(8.5)

        for decl in declarations:
            row_cells = table.add_row().cells
            # Field Name
            f_name = decl.get("field_name") or decl.get("field_type", "").replace("_", " ").title()
            row_cells[0].text = f_name
            row_cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
            row_cells[0].paragraphs[0].runs[0].font.bold = True
            
            # Extracted & Evidence
            val_text = f"\"{decl.get('extracted_text') or '[Field Absent]'}\""
            if decl.get("font_size_display"):
                val_text += f"\n{decl.get('font_size_display')}"
            if decl.get("evidence"):
                val_text += f"\nEvidence: {decl.get('evidence')}"
            row_cells[1].text = val_text
            row_cells[1].paragraphs[0].runs[0].font.size = Pt(8)

            # Detection Rate & Threshold
            conf_val = decl.get("violation_detection_rate")
            if conf_val is None and decl.get("confidence_score") is not None:
                conf_val = int(decl.get("confidence_score", 0.0) * 100)
            thresh_val = decl.get("auto_confirm_threshold") or context.get("auto_confirm_threshold", 80.0)
            verif_method = decl.get("verification_method", "System Auto-Confirmation")
            if conf_val is not None:
                rate_text = f"Confidence: {conf_val}%\nThreshold: {thresh_val}%\n{verif_method}"
            else:
                rate_text = "N/A"
            row_cells[2].text = rate_text
            row_cells[2].paragraphs[0].runs[0].font.size = Pt(7.5)

            # Workflow & Decision
            is_compliant = decl.get("is_compliant", False)
            base_status = "COMPLIANT" if is_compliant else "NON-COMPLIANT"
            wf_dec = decl.get("workflow_decision")
            h_dec = decl.get("human_decision")
            is_officer = bool(h_dec and h_dec not in ["AUTO-ANALYZED", "NONE"])
            if is_officer:
                st_text = f"{base_status}\n[OFFICER {h_dec.upper()}]"
            elif wf_dec in ["AUTO_CONFIRMED", "AUTOMATICALLY_CONFIRMED"] or (conf_val is not None and conf_val >= thresh_val and decl.get("status_label") != "NEEDS MANUAL REVIEW"):
                st_text = f"{base_status}\n[AUTO-CONFIRMED]"
            else:
                st_text = f"{base_status}\n[MANUAL REVIEW REQ]"
            row_cells[3].text = st_text
            row_cells[3].paragraphs[0].runs[0].font.bold = True
            row_cells[3].paragraphs[0].runs[0].font.size = Pt(8)

            # Rule & Reason
            rule_ref = decl.get("rule_reference", "LMPC Rules 2011")
            reason = decl.get("reason") or decl.get("reviewer_notes") or ""
            row_cells[4].text = f"{rule_ref}\n\nReason: {reason}"
            row_cells[4].paragraphs[0].runs[0].font.size = Pt(8)

        # Workflow Disclaimer
        if context.get("workflow_disclaimer"):
            wf_p = doc.add_paragraph()
            wf_run = wf_p.add_run(f"Compliance Workflow Notice: {context.get('workflow_disclaimer')}")
            wf_run.font.size = Pt(8)
            wf_run.font.italic = True
            wf_run.font.color.rgb = RGBColor(0x55, 0x58, 0x69)

        doc.add_paragraph()

        # 3. Statutory Assessment & Overall Result Section
        assess_heading = doc.add_heading(level=2)
        assess_run = assess_heading.add_run("3. Overall Statutory Inspection Result & Verification Basis")
        assess_run.font.size = Pt(12)
        assess_run.font.bold = True
        assess_run.font.color.rgb = RGBColor(0x00, 0x51, 0xD5)

        det_table = doc.add_table(rows=2, cols=2)
        det_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        det_table.cell(0, 0).text = "Overall Inspection Result:"
        det_table.cell(0, 0).paragraphs[0].runs[0].font.bold = True
        det_table.cell(0, 0).paragraphs[0].runs[0].font.size = Pt(8.5)
        det_table.cell(0, 1).text = f"{context.get('overall_label', 'VERIFIED COMPLIANT')}\n(Derived from Verified Declaration Findings)"
        det_table.cell(0, 1).paragraphs[0].runs[0].font.size = Pt(8.5)

        det_table.cell(1, 0).text = "Decision Basis:"
        det_table.cell(1, 0).paragraphs[0].runs[0].font.bold = True
        det_table.cell(1, 0).paragraphs[0].runs[0].font.size = Pt(8.5)
        det_table.cell(1, 1).text = f"{context.get('decision_basis', '')}\nAuto-Confirmed: {context.get('auto_confirmed_count', 0)} | Inspector Reviewed: {context.get('inspector_resolved_count', 0)} | Unresolved: {context.get('manual_review_count', 0)}"
        det_table.cell(1, 1).paragraphs[0].runs[0].font.size = Pt(8.5)

        doc.add_paragraph()

        assess_p = doc.add_paragraph()
        assess_text = context.get("statutory_assessment") or "Inspection findings recorded for regulatory audit."
        assess_p.add_run(f"Statutory Determination: {assess_text}").font.size = Pt(9)

        if context.get("final_remarks") or context.get("custom_notes"):
            notes_p = doc.add_paragraph()
            notes_run = notes_p.add_run(f"Officer Remarks: {context.get('final_remarks') or context.get('custom_notes')}")
            notes_run.font.italic = True
            notes_run.font.size = Pt(9)

        # 4. Audit Trail Table
        audit_trail = context.get("audit_trail", [])
        if audit_trail:
            doc.add_paragraph()
            audit_heading = doc.add_heading(level=2)
            audit_run = audit_heading.add_run("4. Inspection Audit Trail & Manual Review History")
            audit_run.font.size = Pt(12)
            audit_run.font.bold = True
            audit_run.font.color.rgb = RGBColor(0x00, 0x51, 0xD5)

            audit_tbl = doc.add_table(rows=1, cols=4)
            audit_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            a_hdr = audit_tbl.rows[0].cells
            for idx, col_name in enumerate(["Timestamp", "Officer / User", "Action", "Explanation / Remarks"]):
                a_hdr[idx].text = col_name
                a_hdr[idx].paragraphs[0].runs[0].font.bold = True
                a_hdr[idx].paragraphs[0].runs[0].font.size = Pt(8)

            for entry in audit_trail:
                a_row = audit_tbl.add_row().cells
                a_row[0].text = str(entry.get("timestamp", ""))
                a_row[0].paragraphs[0].runs[0].font.size = Pt(7.5)
                a_row[1].text = f"{entry.get('user_name', '')} ({entry.get('user_role', '')})"
                a_row[1].paragraphs[0].runs[0].font.size = Pt(7.5)
                a_row[2].text = str(entry.get("action", ""))
                a_row[2].paragraphs[0].runs[0].font.size = Pt(7.5)
                a_row[3].text = str(entry.get("notes", ""))
                a_row[3].paragraphs[0].runs[0].font.size = Pt(7.5)

        # 5. Signatures
        doc.add_paragraph("\n5. Verification & Authorization\n")
        sig_table = doc.add_table(rows=1, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_cells = sig_table.rows[0].cells
        sig_cells[0].text = f"_____________________________\n{context.get('finalized_by') or context.get('inspector_name')}\nEnforcement Inspector / Legal Metrology Officer\nDepartment of Legal Metrology"
        sig_cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
        sig_cells[1].text = f"_____________________________\nAuthorized Controller of Legal Metrology\nDigital Verification Seal\nDate: {context.get('generated_at')}"
        sig_cells[1].paragraphs[0].runs[0].font.size = Pt(8.5)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

